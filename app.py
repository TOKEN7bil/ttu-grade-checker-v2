from flask import Flask, request, jsonify
from openpyxl import load_workbook

app = Flask(__name__)

# Load Excel with openpyxl
wb = load_workbook('data.xlsx')
sheet = wb.active

# Convert to dictionary
students = {}
headers = [cell.value for cell in sheet[1]]
for row in sheet.iter_rows(min_row=2, values_only=True):
    data = dict(zip(headers, row))
    students[data['IndexNumber']] = data

@app.route('/ussd', methods=['POST'])
def ussd():
    data = request.get_json()
    user_input = data.get('input', '')

    # Step 1: Show menu
    if user_input == '':
        return jsonify({
            "response": "Welcome to TTU Results Checker\n1. Login to view grades\n2. Exit",
            "type": "response"
        })

    # Step 2: User pressed 1
    if user_input == '1':
        return jsonify({
            "response": "Enter IndexNumber*Password",
            "type": "response"
        })

    # Step 3: User entered Index*Pass
    if '*' in user_input:
        try:
            index, password = user_input.split('*')
            if index in students and str(students[index]['Password']) == password:
                grade = students[index]['Grade']
                return jsonify({
                    "response": f"Your Grade: {grade}\nThank you",
                    "type": "end"
                })
            else:
                return jsonify({
                    "response": "Invalid Index or Password",
                    "type": "end"
                })
        except:
            return jsonify({"response": "Wrong format. Use: Index*Password", "type": "end"})

    # Step 2: User pressed 2
    if user_input == '2':
        return jsonify({"response": "Goodbye", "type": "end"})

if __name__ == '__main__':
    app.run()
